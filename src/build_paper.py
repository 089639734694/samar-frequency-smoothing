"""
build_paper.py - Membangun berkas Word makalah dari template IEEE dosen.

Template .doc dikonversi lebih dulu ke .docx (lihat README), lalu isinya
dikosongkan dan diisi ulang memakai style asli template sehingga format
dijamin sama persis dengan yang diminta.

Struktur section:
  seksi 0 : satu kolom  -> blok judul, penulis, afiliasi, abstrak
  seksi 1 : dua kolom   -> isi makalah
  gambar lebar dua kolom disisipkan dengan menutup seksi dua kolom,
  menyelipkan pita satu kolom, lalu kembali ke dua kolom.
"""

from __future__ import annotations

import copy
import os
import re
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

sys.path.insert(0, os.path.dirname(__file__))
import paper_content as C  # noqa: E402

ROOT = os.path.dirname(os.path.dirname(__file__))
TEMPLATE = os.path.join(ROOT, "results", "IEEE_template.docx")
FIGS = os.path.join(ROOT, "results", "figs")
KELUARAN = os.path.join(ROOT, "Makalah-IF6060-23525011-SAMAR.docx")

ROMAWI = ["I", "II", "III", "IV", "V", "VI", "VII", "VIII", "IX", "X", "XI"]


# ---------------------------------------------------------------------------
# Utilitas XML
# ---------------------------------------------------------------------------

def _kosongkan_body(doc):
    """Hapus seluruh isi body, sisakan sectPr terakhir."""
    body = doc.element.body
    sectpr = body.find(qn("w:sectPr"))
    for anak in list(body):
        if anak is not sectpr:
            body.remove(anak)
    return sectpr


def _sectpr_kolom(basis, n_kolom, kontinu=True):
    """Salin sectPr basis lalu setel jumlah kolom dan jenis section."""
    baru = copy.deepcopy(basis)
    for tag in ("w:cols", "w:type"):
        for e in baru.findall(qn(tag)):
            baru.remove(e)
    if kontinu:
        t = OxmlElement("w:type")
        t.set(qn("w:val"), "continuous")
        baru.insert(0, t)
    cols = OxmlElement("w:cols")
    cols.set(qn("w:num"), str(n_kolom))
    cols.set(qn("w:space"), "360")          # jarak antarkolom 0,25 inci
    if n_kolom > 1:
        cols.set(qn("w:equalWidth"), "1")
    baru.append(cols)
    return baru


def _tutup_seksi(par, basis, n_kolom):
    """Lampirkan sectPr ke paragraf sehingga seksi berakhir di situ."""
    pPr = par._p.get_or_add_pPr()
    pPr.append(_sectpr_kolom(basis, n_kolom))


def _buang_penomoran(par):
    """Lepaskan penomoran otomatis dari paragraf.

    Style References pada template dosen sudah membawa daftar bernomor,
    sehingga tanpa ini nomor rujukan tercetak ganda, misalnya "[1] [1]".
    Penomoran dibuat manual agar format kurung sikunya terkendali.
    """
    pPr = par._p.get_or_add_pPr()
    for e in pPr.findall(qn("w:numPr")):
        pPr.remove(e)
    numPr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl"); ilvl.set(qn("w:val"), "0")
    numId = OxmlElement("w:numId"); numId.set(qn("w:val"), "0")
    numPr.append(ilvl); numPr.append(numId)
    pPr.append(numPr)


def _set_layout_fixed(tbl):
    """Kunci layout tabel agar Word menghormati lebar kolom yang disetel."""
    tbl.autofit = False
    tblPr = tbl._tbl.tblPr
    for e in tblPr.findall(qn("w:tblLayout")):
        tblPr.remove(e)
    lay = OxmlElement("w:tblLayout")
    lay.set(qn("w:type"), "fixed")
    tblPr.append(lay)


def _header_berulang(tbl):
    """Tandai baris pertama sebagai header agar berulang bila tabel terpecah."""
    trPr = tbl.rows[0]._tr.get_or_add_trPr()
    e = OxmlElement("w:tblHeader")
    trPr.append(e)


def _jaga_utuh(tbl):
    """Setel keep-with-next pada semua baris kecuali terakhir, agar Word
    memindahkan tabel utuh ke halaman berikutnya alih-alih memecahnya."""
    for row in tbl.rows[:-1]:
        for sel in row.cells:
            for par in sel.paragraphs:
                par.paragraph_format.keep_with_next = True
    for row in tbl.rows:
        trPr = row._tr.get_or_add_trPr()
        trPr.append(OxmlElement("w:cantSplit"))


def _set_border_tabel(tbl):
    tblPr = tbl._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for sisi in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{sisi}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "4")
        e.set(qn("w:color"), "000000")
        borders.append(e)
    tblPr.append(borders)


# ---------------------------------------------------------------------------
# Penyisip konten
# ---------------------------------------------------------------------------

class Pembangun:
    def __init__(self):
        self.doc = Document(TEMPLATE)
        # Ketentuan tugas: Times New Roman 10, spasi 1. Template dosen
        # memakai spasi 1,05 pada style isi, jadi disetel ulang ke 1,0.
        for nm in ("Text", "Equation", "Abstract", "References",
                   "Figure Caption"):
            try:
                st = self.doc.styles[nm]
                st.paragraph_format.line_spacing = 1.0
            except KeyError:
                pass
        self.sectpr_akhir = _kosongkan_body(self.doc)
        self.basis = copy.deepcopy(self.sectpr_akhir)
        self.no_h1 = 0
        self.no_h2 = 0
        # pastikan seksi terakhir dua kolom
        for tag in ("w:cols",):
            for e in self.sectpr_akhir.findall(qn(tag)):
                self.sectpr_akhir.remove(e)
        cols = OxmlElement("w:cols")
        cols.set(qn("w:num"), "2")
        cols.set(qn("w:space"), "360")
        cols.set(qn("w:equalWidth"), "1")
        self.sectpr_akhir.append(cols)

    # -- primitif -------------------------------------------------------
    def p(self, teks, style="Text", align=None, size=None, italic=None,
          bold=None, space_after=None):
        par = self.doc.add_paragraph(style=style)
        run = par.add_run(teks)
        if size is not None:
            run.font.size = Pt(size)
        if italic is not None:
            run.font.italic = italic
        if bold is not None:
            run.font.bold = bold
        if align is not None:
            par.alignment = align
        if space_after is not None:
            par.paragraph_format.space_after = Pt(space_after)
        return par

    def h1(self, teks):
        self.no_h1 += 1
        self.no_h2 = 0
        return self.p(f"{ROMAWI[self.no_h1 - 1]}.  {teks}", style="Heading 1")

    def h2(self, teks):
        self.no_h2 += 1
        huruf = chr(ord("A") + self.no_h2 - 1)
        return self.p(f"{huruf}. {teks}", style="Heading 2")

    def daftar(self, butir):
        for i, b in enumerate(butir, 1):
            par = self.p(f"{i})  {b}", style="Text")
            par.paragraph_format.left_indent = Inches(0.18)
            par.paragraph_format.first_line_indent = Inches(-0.18)

    def tabel(self, kunci, nomor):
        spec = C.TABEL[kunci]
        jn = self.p(f"TABEL {nomor}", style="Text",
                    align=WD_ALIGN_PARAGRAPH.CENTER, size=8, space_after=0)
        jn.paragraph_format.keep_with_next = True
        jt = self.p(spec["judul"], style="Text",
                    align=WD_ALIGN_PARAGRAPH.CENTER, size=8, space_after=3)
        jt.paragraph_format.keep_with_next = True

        n_kol = len(spec["header"])
        tbl = self.doc.add_table(rows=1, cols=n_kol)
        _set_border_tabel(tbl)
        _set_layout_fixed(tbl)
        hdr = tbl.rows[0].cells
        for i, h in enumerate(spec["header"]):
            hdr[i].text = ""
            par = hdr[i].paragraphs[0]
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = par.add_run(h)
            r.font.size = Pt(7)
            r.font.bold = True
        for baris in spec["baris"]:
            sel = tbl.add_row().cells
            for i, v in enumerate(baris):
                sel[i].text = ""
                par = sel[i].paragraphs[0]
                if i > 0:
                    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r = par.add_run(v)
                r.font.size = Pt(7)
                # baris sub-judul di dalam tabel dicetak tebal
                if i == 0 and v and not v.startswith(" ") and not baris[1]:
                    r.font.bold = True
        _header_berulang(tbl)
        _jaga_utuh(tbl)
        for i, w in enumerate(spec.get("lebar", [])):
            tbl.columns[i].width = Inches(w)
            for row in tbl.rows:
                row.cells[i].width = Inches(w)
        if spec.get("catatan"):
            self.p(spec["catatan"], style="Text", size=7, italic=True)
        self.p("", style="Text", size=4, space_after=0)

    def gambar(self, kunci):
        g = C.GAMBAR[kunci]
        par = self.doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.space_after = Pt(2)
        par.add_run().add_picture(os.path.join(FIGS, g["berkas"]),
                                  width=Inches(g["lebar_in"]))
        cap = self.p(g["caption"], style="Figure Caption")
        cap.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        self.p("", style="Text", size=4, space_after=0)

    # -- perakitan ------------------------------------------------------
    def blok_judul(self):
        self.p(C.JUDUL, style="Paper Title")
        self.p(C.PENULIS, style="Authors")
        for a in C.AFILIASI:
            self.p(a, style="Affiliations")
        self.p(C.EMAIL, style="Affiliations")
        self.p("", style="Text", size=6, space_after=0)
        self.p(C.ABSTRAK, style="Abstract")
        akhir = self.p(C.KATA_KUNCI, style="Abstract")
        # tutup seksi satu kolom di sini
        _tutup_seksi(akhir, self.basis, 1)

    def isi(self):
        no_tabel = 0
        romawi_tabel = ["I", "II", "III", "IV", "V", "VI", "VII", "VIII"]
        for jenis, isi in C.ISI:
            if jenis == "H1":
                self.h1(isi)
            elif jenis == "H2":
                self.h2(isi)
            elif jenis == "P":
                self.p(isi)
            elif jenis == "EQ":
                self.p(isi, style="Equation",
                       align=WD_ALIGN_PARAGRAPH.CENTER)
            elif jenis == "LIST":
                self.daftar(isi)
            elif jenis == "TBL":
                no_tabel += 1
                if C.TABEL[isi].get("dua_kolom"):
                    self._pita_satu_kolom(
                        lambda: self.tabel(isi, romawi_tabel[no_tabel - 1]))
                else:
                    self.tabel(isi, romawi_tabel[no_tabel - 1])
            elif jenis == "FIG":
                if C.GAMBAR[isi]["dua_kolom"]:
                    self._gambar_lebar(isi)
                else:
                    self.gambar(isi)

    def _pita_satu_kolom(self, isi_fn):
        """Buka pita selebar halaman, jalankan isi_fn, lalu kembali 2 kolom."""
        penutup = self.p("", style="Text", size=2, space_after=0)
        _tutup_seksi(penutup, self.basis, 2)
        isi_fn()
        penutup2 = self.p("", style="Text", size=2, space_after=0)
        _tutup_seksi(penutup2, self.basis, 1)

    def _gambar_lebar(self, kunci):
        """Sisipkan pita satu kolom untuk gambar selebar dua kolom."""
        self._pita_satu_kolom(lambda: self.gambar(kunci))

    def penutup(self):
        self.p("REFERENSI", style="Reference Head",
               align=WD_ALIGN_PARAGRAPH.CENTER)
        for i, r in enumerate(C.REFERENSI, 1):
            par = self.p(f"[{i}]  {r}", style="References")
            _buang_penomoran(par)
            par.paragraph_format.left_indent = Inches(0.22)
            par.paragraph_format.first_line_indent = Inches(-0.22)
        self.p("", style="Text", size=6, space_after=0)
        self.p("PERNYATAAN", style="Reference Head",
               align=WD_ALIGN_PARAGRAPH.CENTER)
        self.p(C.PERNYATAAN, style="Text")
        self.p("", style="Text", size=6, space_after=0)
        self.p(C.TANGGAL, style="Text")
        self.p(C.TTD, style="Text")

    def simpan(self):
        self.doc.save(KELUARAN)
        return KELUARAN


def main():
    b = Pembangun()
    b.blok_judul()
    b.isi()
    b.penutup()
    path = b.simpan()
    print("Makalah ditulis ke:", path)
    return path


if __name__ == "__main__":
    main()
